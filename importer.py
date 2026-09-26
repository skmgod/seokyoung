"""기존 프로그램에서 내보낸 엑셀(빌라·세대·미납) 가져오기."""
import csv
import datetime as dt
import io
import re

from openpyxl import Workbook, load_workbook

from billing import reallocate_unit

FIELDS = [
    ("building", "빌라명", ["빌라명", "건물명", "빌라", "빌라명칭", "거래처명", "단지명"]),
    ("address", "주소", ["주소", "신주소", "신)주소", "빌라주소"]),
    ("ho", "호수", ["호수", "호", "동호", "동/호", "동호수"]),
    ("tenant", "성명", ["성명", "입주자", "계약자", "세대주", "이름", "고객명", "입주자명"]),
    ("mobile", "핸드폰", ["핸드폰", "휴대폰", "휴대전화", "핸드폰번호"]),
    ("phone", "전화번호", ["전화번호", "집전화", "연락처", "전화"]),
    ("move_in", "입주일", ["입주일", "입주일자"]),
    ("deposit", "보증금", ["보증금"]),
    ("rent", "임대료", ["임대료", "월세"]),
    ("mgmt", "월관리비", ["월관리비", "관리비", "일반관리비"]),
    ("unpaid", "미납액", ["미납액", "미납금액", "미납", "체납액", "미납합계"]),
    ("status", "상태", ["상태", "입주구분", "공실"]),
    ("memo", "비고", ["비고", "적요", "메모"]),
]
LABELS = {k: label for k, label, _ in FIELDS}


def template_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "세대목록"
    ws.append([label for _, label, _ in FIELDS])
    # 예시는 따로 시트에 둔다(첫 시트만 읽으므로 예시가 실제 자료로 들어가지 않음)
    ex = wb.create_sheet("작성예시")
    ex.append([label for _, label, _ in FIELDS])
    ex.append(["예시빌라", "○○시 ○○구 ○○로 12", "101", "홍길동", "010-0000-0000", "", "2024-03-01",
               5000000, 400000, 50000, 0, "입주", ""])
    ex.append(["예시빌라", "○○시 ○○구 ○○로 12", "102", "", "", "", "", 0, 0, 50000, 0, "공실", ""])
    for sheet in (ws, ex):
        for col, width in zip("ABCDEFGHIJKLM", [14, 28, 8, 10, 15, 14, 12, 12, 10, 10, 10, 8, 20]):
            sheet.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _norm(s):
    return re.sub(r"\s+", "", str(s or ""))


def _read_rows(filename, data):
    if filename.lower().endswith(".csv"):
        for enc in ("utf-8-sig", "cp949"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        return [row for row in csv.reader(io.StringIO(text))]
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _map_header(header):
    cols = [_norm(h) for h in header]
    mapping = {}
    for key, _, aliases in FIELDS:
        for alias in aliases:  # 정확히 같은 이름 우선
            if alias in cols and cols.index(alias) not in mapping.values():
                mapping[key] = cols.index(alias)
                break
        else:
            for i, c in enumerate(cols):
                if i not in mapping.values() and any(c.startswith(a) for a in aliases):
                    mapping[key] = i
                    break
    return mapping


def _to_int(v):
    if v is None or v == "":
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    s = re.sub(r"[^\d\-]", "", str(v))
    return int(s) if s not in ("", "-") else 0


def _to_date(v):
    if v is None or v == "":
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    digits = re.sub(r"\D", "", str(v))
    if len(digits) == 8:
        return f"{digits[:4]}-{digits[4:6]}-{digits[6:]}"
    return ""


def _to_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse(filename, data):
    rows = _read_rows(filename, data)
    header_idx = None
    for i, row in enumerate(rows[:15]):
        m = _map_header(row)
        if "ho" in m and ("building" in m or "tenant" in m):
            header_idx, mapping = i, m
            break
    if header_idx is None:
        raise ValueError("머리글 줄을 찾지 못했습니다. 첫 줄에 '빌라명', '호수' 같은 열 이름이 있어야 합니다.")

    parsed, errors = [], []
    last_building = ""
    for n, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        get = lambda k: row[mapping[k]] if k in mapping and mapping[k] < len(row) else None
        if all(c in (None, "") for c in row):
            continue
        rec = {
            "building": _to_str(get("building")) or last_building,
            "address": _to_str(get("address")),
            "ho": _to_str(get("ho")),
            "tenant": _to_str(get("tenant")),
            "mobile": _to_str(get("mobile")),
            "phone": _to_str(get("phone")),
            "move_in": _to_date(get("move_in")),
            "deposit": _to_int(get("deposit")),
            "rent": _to_int(get("rent")),
            "mgmt": _to_int(get("mgmt")),
            "unpaid": _to_int(get("unpaid")),
            "status": "공실" if "공실" in _to_str(get("status")) else "입주",
            "memo": _to_str(get("memo")),
            "row": n,
        }
        last_building = rec["building"]
        if not rec["building"] or not rec["ho"]:
            errors.append(f"{n}번째 줄: 빌라명 또는 호수가 없어 건너뜁니다")
            continue
        parsed.append(rec)
    return {"mapping": {LABELS[k]: _to_str(rows[header_idx][i]) for k, i in mapping.items()},
            "rows": parsed, "errors": errors}


def commit(conn, rows, carry_ym, carry_date):
    stats = {"buildings": 0, "units_new": 0, "units_updated": 0, "carry": 0}
    mgmt_item = conn.execute("SELECT id FROM fee_items WHERE name = '일반관리비'").fetchone()
    if mgmt_item is None:
        mgmt_item_id = conn.execute("INSERT INTO fee_items(name, sort) VALUES ('일반관리비', 0)").lastrowid
    else:
        mgmt_item_id = mgmt_item["id"]

    for r in rows:
        b = conn.execute("SELECT * FROM buildings WHERE name = ?", (r["building"],)).fetchone()
        if b is None:
            bid = conn.execute("INSERT INTO buildings(name, address) VALUES (?, ?)", (r["building"], r["address"])).lastrowid
            stats["buildings"] += 1
        else:
            bid = b["id"]
            if r["address"] and not b["address"]:
                conn.execute("UPDATE buildings SET address = ? WHERE id = ?", (r["address"], bid))

        u = conn.execute("SELECT id FROM units WHERE building_id = ? AND ho = ?", (bid, r["ho"])).fetchone()
        vals = (r["tenant"], r["phone"], r["mobile"], r["move_in"], r["status"], r["deposit"], r["rent"], r["memo"])
        if u is None:
            uid = conn.execute(
                """INSERT INTO units(tenant, phone, mobile, move_in, status, deposit, rent, memo, building_id, ho)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", vals + (bid, r["ho"])).lastrowid
            stats["units_new"] += 1
        else:
            uid = u["id"]
            conn.execute(
                """UPDATE units SET tenant = ?, phone = ?, mobile = ?, move_in = ?, status = ?, deposit = ?, rent = ?,
                   memo = ? WHERE id = ?""", vals + (uid,))
            stats["units_updated"] += 1

        if r["mgmt"]:
            conn.execute("INSERT OR IGNORE INTO building_fees(building_id, item_id, amount, method) VALUES (?, ?, 0, 'unit')",
                         (bid, mgmt_item_id))
            conn.execute("INSERT OR REPLACE INTO unit_fees(unit_id, item_id, amount) VALUES (?, ?, ?)",
                         (uid, mgmt_item_id, r["mgmt"]))

        if r["unpaid"]:
            conn.execute("DELETE FROM bills WHERE unit_id = ? AND kind = 'carry'", (uid,))
            bill_id = conn.execute(
                "INSERT INTO bills(unit_id, ym, kind, amount, late_fee, due_date, issued, memo) VALUES (?, ?, 'carry', ?, 0, ?, ?, '기존 프로그램 이월')",
                (uid, carry_ym, r["unpaid"], carry_date, carry_date)).lastrowid
            conn.execute("INSERT INTO bill_lines(bill_id, name, amount, detail) VALUES (?, '이월 미납액', ?, '기존 프로그램에서 이월')",
                         (bill_id, r["unpaid"]))
            reallocate_unit(conn, uid)
            stats["carry"] += 1
    return stats
